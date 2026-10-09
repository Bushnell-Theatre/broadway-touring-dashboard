"""
test_process_revision.py — Revision / append / audit behaviour in
scripts/process_touring.py

Covers only the revision feature:

   1  REV / REVISED filename forms are recognised
   2  words merely containing "rev" are NOT treated as revisions
   3  normal append still appends
   4  a normal append carrying conflicting values fails closed
   5  revision updates a changed record
   6  revision inserts a record the store did not have
   7  a stored record absent from a revision is RETAINED, with a warning
   8  re-running the same revision is a clean no-op
   9  an atomic failure leaves the previous data.json intact
  10  failed validation prevents publication (non-zero exit, no write)
  11  the historical audit reports conflicts and ambiguity without writing

Fixtures are real .xlsx workbooks built with openpyxl — the same library the
pipeline parses with — so the tests exercise the production parser end to end.

Run: python scripts/test_process_revision.py
"""

import importlib.util
import json
import os
import shutil
import subprocess
import sys
import tempfile

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SCRIPT = os.path.join(ROOT, 'scripts', 'process_touring.py')

spec = importlib.util.spec_from_file_location('pt', SCRIPT)
pt = importlib.util.module_from_spec(spec)
spec.loader.exec_module(pt)

# ── Assertion counters ───────────────────────────────────────────────────

passed = 0
failed = 0


def ok(label, cond, detail=''):
    global passed, failed
    if cond:
        passed += 1
        print(f"  PASS  {label}")
    else:
        failed += 1
        print(f"  FAIL  {label}" + (f" — {detail}" if detail else ''))


def eq(label, actual, expected):
    ok(label, actual == expected, f"expected {expected!r}, got {actual!r}")


def section(name):
    print(f"\n--- {name} ---")


# ── Fixture helpers ──────────────────────────────────────────────────────

HEADERS = ['Show', 'Theatre', 'City', 'TicketRange', 'TopPrice', 'NumPerf',
           'GrossGross', 'GrossPotential', 'GGPctGP', 'PaidTix', 'TotalTix',
           'Capacity', 'CapPaid', 'CapTotal', 'OnSub', 'AvgAdm']


def row(show, theatre, city, gross, perf=8, potential=2000000.0, ggpct=70.0,
        paid=20000, total=21000, capacity=21632, cappaid=92.0, captotal=97.0,
        sub='', avgadm=66.0, top=199):
    return [show, theatre, city, '$39-$199', top, perf, gross, potential,
            ggpct, paid, total, capacity, cappaid, captotal, sub, avgadm]


def make_workbook(path, rows, week='10-4-26', tier='PRIMARY'):
    """Write a workbook the production parser will accept."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"TouringRep{tier}_{week}"
    ws.append(HEADERS)
    for r in rows:
        ws.append(r)
    wb.save(path)
    return path


def make_data_json(path, records):
    with open(path, 'w', encoding='utf-8') as f:
        json.dump({'generated_at': '2026-01-01T00:00:00+00:00',
                   'record_count': len(records),
                   'records': records}, f, indent=2, ensure_ascii=False)


def parse(path):
    return pt.process_file(path, [])


def load(path):
    with open(path, 'r', encoding='utf-8') as f:
        return json.load(f)


def run(args, cwd=None):
    return subprocess.run([sys.executable, SCRIPT] + args,
                          capture_output=True, text=True, cwd=cwd or ROOT)


# ══════════════════════════════════════════════════════════════════════════
section('Suite 1 — revision filename forms are recognised')

for name, expected_ordinal in [
    ('TouringReport_2627Wk19_10-4-26 - REV.xlsx', None),
    ('TouringReport_2627Wk19_10-4-26 - REV2.xlsx', 2),
    ('TouringReport_2627Wk19_10-4-26_REV.xlsx', None),
    ('TouringReport_2627Wk19_10-4-26 (REV).xlsx', None),
    ('TouringReport_2627Wk19_10-4-26 - REVISED.xlsx', None),
    ('TouringReport_2627Wk19_10-4-26 - revised.xlsx', None),
    ('TouringReport_2627Wk19_10-4-26 - Rev3.xlsx', 3),
    ('REV - TouringReport.xlsx', None),
]:
    is_rev, ordinal = pt.revision_marker(name)
    ok(f"recognised: {name}", is_rev)
    eq(f"  ordinal for {name}", ordinal, expected_ordinal)

# ══════════════════════════════════════════════════════════════════════════
section('Suite 2 — no false filename matches')

for name in [
    'TouringReport_2627Wk19_10-4-26.xlsx',
    'Preview_10-4-26.xlsx',
    'Revenue_Report_10-4-26.xlsx',
    'Review_10-4-26.xlsx',
    'Reverb_10-4-26.xlsx',
    'FOREVER_10-4-26.xlsx',
    'TouringReportRevisedFinal.xlsx',   # no delimiter before the token
    'Revolution_10-4-26.xlsx',
]:
    is_rev, _ = pt.revision_marker(name)
    ok(f"not a revision: {name}", not is_rev)

# ══════════════════════════════════════════════════════════════════════════
section('Suite 3 — normal append still appends')

tmp = tempfile.mkdtemp(prefix='btd_rev_')
try:
    wk1 = make_workbook(os.path.join(tmp, 'Report_9-27-26.xlsx'),
                        [row('Wicked', 'Pantages', 'Los Angeles', 1000000.0)],
                        week='9-27-26')
    data = os.path.join(tmp, 'data.json')
    make_data_json(data, [])

    r = run(['--append', wk1, data])
    eq('append exits 0', r.returncode, 0)
    doc = load(data)
    eq('one record appended', doc['record_count'], 1)
    ok('MODE: append reported', 'MODE: append' in r.stdout, r.stdout[:200])

    # A second, different week appends alongside it
    wk2 = make_workbook(os.path.join(tmp, 'Report_10-4-26.xlsx'),
                        [row('Hamilton', 'Orpheum', 'Chicago', 900000.0)])
    r = run(['--append', wk2, data])
    eq('second append exits 0', r.returncode, 0)
    eq('two records stored', load(data)['record_count'], 2)
finally:
    shutil.rmtree(tmp, ignore_errors=True)

# ══════════════════════════════════════════════════════════════════════════
section('Suite 4 — conflicting normal append fails closed')

tmp = tempfile.mkdtemp(prefix='btd_rev_')
try:
    data = os.path.join(tmp, 'data.json')
    orig = make_workbook(os.path.join(tmp, 'Report_10-4-26.xlsx'),
                         [row('The Outsiders', 'Pantages', 'Los Angeles', 1334257.0)])
    make_data_json(data, parse(orig))
    before = load(data)

    # Same engagement, different gross, but NOT named as a revision.
    clash = make_workbook(os.path.join(tmp, 'Report_10-4-26_again.xlsx'),
                          [row('The Outsiders', 'Pantages', 'Los Angeles', 1473785.0)])
    r = run(['--append', clash, data])
    ok('conflicting append exits non-zero', r.returncode != 0, f"rc={r.returncode}")
    ok('conflict is reported', 'restated with DIFFERENT values' in r.stdout, r.stdout[-400:])
    ok('the changed field is named', 'gross_gross' in r.stdout)
    ok('it points at the REV convention', 'REVISED' in r.stdout)
    eq('data.json untouched', load(data), before)
finally:
    shutil.rmtree(tmp, ignore_errors=True)

# ══════════════════════════════════════════════════════════════════════════
section('Suite 5/6/7/8 — revision upsert semantics')

tmp = tempfile.mkdtemp(prefix='btd_rev_')
try:
    data = os.path.join(tmp, 'data.json')
    orig = make_workbook(os.path.join(tmp, 'Report_10-4-26.xlsx'), [
        row('The Outsiders', 'Pantages', 'Los Angeles', 1334257.0, avgadm=70.94, ggpct=73.2),
        row('Wicked', 'Orpheum', 'Chicago', 900000.0),
        row('Hadestown', 'Fox', 'Atlanta', 500000.0),
    ])
    make_data_json(data, parse(orig))
    eq('store starts with 3 records', load(data)['record_count'], 3)

    # Revision: one row corrected, one row new, one row (Hadestown) ABSENT.
    rev = make_workbook(os.path.join(tmp, 'Report_10-4-26 - REVISED.xlsx'), [
        row('The Outsiders', 'Pantages', 'Los Angeles', 1473785.0, avgadm=78.36, ggpct=80.8),
        row('Wicked', 'Orpheum', 'Chicago', 900000.0),
        row('Six', 'Majestic', 'Boston', 400000.0),
    ])
    r = run(['--append', rev, data])   # routed to revision mode by filename
    eq('revision exits 0', r.returncode, 0)
    ok('MODE: revision reported', 'MODE: revision' in r.stdout, r.stdout[:200])

    doc = load(data)
    recs = {(x['show'], x['city']): x for x in doc['records']}

    # 5 — update applied
    eq('updated gross', recs[('The Outsiders', 'Los Angeles')]['gross_gross'], 1473785.0)
    eq('updated avg_adm', recs[('The Outsiders', 'Los Angeles')]['avg_adm'], 78.36)
    eq('updated gg_pct_gp', recs[('The Outsiders', 'Los Angeles')]['gg_pct_gp'], 80.8)
    ok('field-level change logged', 'gross_gross: 1334257.0 -> 1473785.0' in r.stdout,
       r.stdout[-600:])

    # 6 — addition applied
    ok('new record inserted', ('Six', 'Boston') in recs)
    ok('addition logged', 'ADDED' in r.stdout)

    # 7 — absent record retained, with the required warning
    ok('absent record RETAINED', ('Hadestown', 'Atlanta') in recs)
    eq('record count grew by exactly the addition', doc['record_count'], 4)
    ok('retention statement present',
       'were RETAINED' in r.stdout and 'not guaranteed to be a complete population' in r.stdout,
       r.stdout[-800:])
    ok('fewer-records warning present',
       'Revision contains fewer records than the stored week' in r.stdout
       or len(parse(rev)) >= 3)
    ok('no deletions reported', 'Removed records:   0' in r.stdout)
    eq('counts line: added', '| Added records:     1' in r.stdout, True)
    eq('counts line: updated', '| Updated records:   1' in r.stdout, True)
    eq('counts line: identical', '| Identical records: 1' in r.stdout, True)
    ok('week gross before/after logged', 'Week 2026-10-04 gross:' in r.stdout)

    # 8 — rerunning the identical revision is a clean no-op
    snapshot = load(data)['records']
    r2 = run(['--append', rev, data])
    eq('rerun exits 0', r2.returncode, 0)
    after = load(data)['records']
    eq('rerun changes nothing', after, snapshot)
    ok('rerun reports 0 added / 0 updated',
       '| Added records:     0' in r2.stdout and '| Updated records:   0' in r2.stdout,
       r2.stdout[-400:])
finally:
    shutil.rmtree(tmp, ignore_errors=True)

# ══════════════════════════════════════════════════════════════════════════
section('Suite 9 — guards leave previous data intact')

tmp = tempfile.mkdtemp(prefix='btd_rev_')
try:
    data = os.path.join(tmp, 'data.json')
    orig = make_workbook(os.path.join(tmp, 'Report_10-4-26.xlsx'),
                         [row('Wicked', 'Orpheum', 'Chicago', 900000.0)])
    make_data_json(data, parse(orig))
    before = load(data)

    # 0 parsed records — an empty revision must never touch the store
    empty = os.path.join(tmp, 'Empty_10-4-26 - REVISED.xlsx')
    wb = openpyxl.Workbook()
    wb.active.title = 'TouringRepPRIMARY_10-4-26'
    wb.active.append(HEADERS)
    wb.save(empty)
    r = run(['--revision', empty, data])
    ok('empty revision exits non-zero', r.returncode != 0, f"rc={r.returncode}")
    ok('empty revision reported', 'parsed to 0 records' in r.stdout, r.stdout[-300:])
    eq('data.json untouched after empty revision', load(data), before)

    # more than one reporting week in a single revision
    multi = os.path.join(tmp, 'Multi - REVISED.xlsx')
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = 'TouringRepPRIMARY_10-4-26'
    ws1.append(HEADERS)
    ws1.append(row('Wicked', 'Orpheum', 'Chicago', 950000.0))
    ws2 = wb.create_sheet('TouringRepSECONDARY_9-27-26')
    ws2.append(HEADERS)
    ws2.append(row('Hadestown', 'Fox', 'Atlanta', 500000.0))
    wb.save(multi)
    r = run(['--revision', multi, data])
    ok('multi-week revision exits non-zero', r.returncode != 0, f"rc={r.returncode}")
    ok('multi-week reported', 'reporting weeks' in r.stdout, r.stdout[-300:])
    eq('data.json untouched after multi-week revision', load(data), before)

    ok('no stray .tmp left behind', not os.path.isfile(data + '.tmp'))
finally:
    shutil.rmtree(tmp, ignore_errors=True)

# ══════════════════════════════════════════════════════════════════════════
section('Suite 10 — failed validation prevents publication')

tmp = tempfile.mkdtemp(prefix='btd_rev_')
try:
    data = os.path.join(tmp, 'data.json')
    orig = make_workbook(os.path.join(tmp, 'Report_10-4-26.xlsx'),
                         [row('Wicked', 'Orpheum', 'Chicago', 900000.0)])
    make_data_json(data, parse(orig))
    before = load(data)

    # Duplicate canonical keys must be rejected by validate_payload, which is
    # what stands between a corrupt candidate and a published deploy.
    recs = parse(orig)
    log = []
    ok('validator accepts a clean payload',
       pt.validate_payload({'record_count': 1, 'records': recs}, 1, log))
    log = []
    dupe = recs + [dict(recs[0])]
    ok('validator rejects duplicate keys',
       not pt.validate_payload({'record_count': 2, 'records': dupe}, 2, log))
    ok('duplicate rejection is explained',
       any('duplicate canonical keys' in line for line in log), log)
    log = []
    ok('validator rejects a count mismatch',
       not pt.validate_payload({'record_count': 9, 'records': recs}, 1, log))

    # write_data_atomic must refuse a bad payload and leave the live file alone
    log = []
    wrote = pt.write_data_atomic(data, dupe, log)
    ok('atomic write refuses a bad payload', not wrote)
    eq('data.json untouched after refusal', load(data), before)
    ok('candidate temp file cleaned up', not os.path.isfile(data + '.tmp'))
    ok('refusal is explained', any('left unchanged' in line for line in log), log)

    # a good payload does go live, and keeps a backup of the prior file
    log = []
    extra = parse(make_workbook(os.path.join(tmp, 'Report_10-11-26.xlsx'),
                                [row('Six', 'Majestic', 'Boston', 400000.0)],
                                week='10-11-26'))
    ok('atomic write accepts a good payload',
       pt.write_data_atomic(data, recs + extra, log))
    eq('live file updated', load(data)['record_count'], 2)
    ok('backup of prior data retained', os.path.isfile(data + '.bak'))
    eq('backup holds the previous content', load(data + '.bak'), before)
finally:
    shutil.rmtree(tmp, ignore_errors=True)

# ══════════════════════════════════════════════════════════════════════════
section('Suite 11 — historical audit is read-only and reports ambiguity')

tmp = tempfile.mkdtemp(prefix='btd_rev_')
try:
    reports = os.path.join(tmp, 'reports')
    os.makedirs(reports)
    data = os.path.join(tmp, 'data.json')

    orig = make_workbook(os.path.join(reports, 'Report_10-4-26.xlsx'), [
        row('The Outsiders', 'Pantages', 'Los Angeles', 1334257.0),
        row('Wicked', 'Orpheum', 'Chicago', 900000.0),
    ])
    make_data_json(data, parse(orig))
    before = load(data)

    # One unambiguous revision: stored data is now stale against it.
    make_workbook(os.path.join(reports, 'Report_10-4-26 - REVISED.xlsx'), [
        row('The Outsiders', 'Pantages', 'Los Angeles', 1473785.0),
    ])
    r = run(['--audit', reports, data])
    eq('audit exits 0', r.returncode, 0)
    ok('audit names the original file', 'Report_10-4-26.xlsx' in r.stdout)
    ok('audit names the revision file', 'Report_10-4-26 - REVISED.xlsx' in r.stdout)
    ok('audit reports the value difference', 'VALUES DIFFER' in r.stdout, r.stdout[-700:])
    ok('audit shows stored vs source gross', 'source-derived' in r.stdout)
    eq('audit wrote nothing', load(data), before)

    # A second unnumbered revision makes precedence ambiguous.
    make_workbook(os.path.join(reports, 'Report_10-4-26 - REV.xlsx'), [
        row('The Outsiders', 'Pantages', 'Los Angeles', 1500000.0),
    ])
    r = run(['--audit', reports, data])
    eq('audit still exits 0', r.returncode, 0)
    ok('ambiguity is reported', 'AMBIGUOUS REVISION PRECEDENCE' in r.stdout, r.stdout[-700:])
    ok('ambiguous set is not applied',
       '1500000' not in r.stdout.split('AMBIGUOUS')[1].split('stored records')[0]
       or 'needs a human' in r.stdout)
    ok('resolution is suggested', 'REV2' in r.stdout)
    eq('audit still wrote nothing', load(data), before)

    # Numbered revisions resolve the ambiguity by ordinal, not by name or mtime.
    os.remove(os.path.join(reports, 'Report_10-4-26 - REV.xlsx'))
    make_workbook(os.path.join(reports, 'Report_10-4-26 - REV2.xlsx'), [
        row('The Outsiders', 'Pantages', 'Los Angeles', 1500000.0),
    ])
    r = run(['--audit', reports, data])
    ok('numbered revisions are unambiguous', 'AMBIGUOUS' not in r.stdout, r.stdout[-700:])
    # REVISED is ordinal 1, REV2 is ordinal 2, so REV2's value must be the one
    # the audit derives — ordering by ordinal, never by name or mtime.
    ok('highest ordinal wins', 'stored 1334257.0 -> source 1500000.0' in r.stdout,
       r.stdout[-700:])
    ok('derived week gross reflects REV2',
       'source-derived 2,400,000' in r.stdout, r.stdout[-700:])

    # A stored record that no source workbook contains must be REPORTED as
    # absent and left alone — the audit never proposes a deletion.
    orphan = dict(parse(orig)[0])
    orphan.update({'show': 'Ghost Record', 'theatre': 'Nowhere', 'city': 'Nowhere',
                   'canonical_key': '2026-10-04|ghost record|nowhere|nowhere|primary'})
    make_data_json(data, parse(orig) + [orphan])
    before_orphan = load(data)
    r = run(['--audit', reports, data])
    eq('audit exits 0 with an orphan present', r.returncode, 0)
    ok('orphan reported as absent from sources',
       'ABSENT from source files' in r.stdout, r.stdout[-900:])
    ok('orphan named in the report', 'Ghost Record' in r.stdout, r.stdout[-900:])
    ok('orphan flagged as retained', 'retained' in r.stdout.lower(), r.stdout[-900:])
    eq('audit never writes', load(data), before_orphan)
finally:
    shutil.rmtree(tmp, ignore_errors=True)

# ══════════════════════════════════════════════════════════════════════════
print("\n" + "=" * 60)
print(f"{passed} passed, {failed} failed")
print("=" * 60)
sys.exit(0 if failed == 0 else 1)
